"""
JSON and Excel Exporters.
Section 19: Full canonical JSON, XLSX with multiple sheets.
"""
from typing import List, Dict, Any
import json
import io
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment


def export_invoices_to_json(invoices: List[Dict[str, Any]]) -> str:
    """Export complete invoice structures to indented JSON string."""
    return json.dumps(invoices, indent=2, default=str)


def export_invoices_to_xlsx(invoices: List[Dict[str, Any]]) -> bytes:
    """
    Export invoice data to multi-sheet Excel workbook:
    - Sheet 1: Bills summary
    - Sheet 2: Line Items
    - Sheet 3: Validation Report
    """
    wb = openpyxl.Workbook()
    # Sheet 1: Bills
    ws_bills = wb.active
    ws_bills.title = "Bills"
    bill_headers = ["Bill Number", "Supplier", "Buyer", "Date", "Due Date", "Subtotal", "Tax", "Total", "Status", "Payment"]
    ws_bills.append(bill_headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")

    for col in range(1, len(bill_headers) + 1):
        ws_bills.cell(row=1, column=col).font = header_font
        ws_bills.cell(row=1, column=col).fill = header_fill

    for inv in invoices:
        ws_bills.append([
            inv.get("bill_number", ""),
            inv.get("supplier_name", ""),
            inv.get("buyer_name", ""),
            inv.get("invoice_date", ""),
            inv.get("due_date", ""),
            float(inv.get("subtotal") or 0.0),
            float(inv.get("tax_amount") or 0.0),
            float(inv.get("total_amount") or 0.0),
            inv.get("validation_status", "valid"),
            inv.get("payment_status", "unpaid")
        ])

    # Sheet 2: Line Items
    ws_items = wb.create_sheet(title="Line Items")
    item_headers = ["Bill Number", "Description", "HSN/SAC", "Quantity", "Unit Price", "Taxable", "Total"]
    ws_items.append(item_headers)
    for col in range(1, len(item_headers) + 1):
        ws_items.cell(row=1, column=col).font = header_font
        ws_items.cell(row=1, column=col).fill = header_fill

    for inv in invoices:
        can = inv.get("canonical_json", {})
        items = can.get("line_items", []) or can.get("items", [])
        for it in items:
            ws_items.append([
                inv.get("bill_number", ""),
                it.get("description", ""),
                it.get("hsn_code") or it.get("hsn_sac") or "",
                float(it.get("quantity") or 1),
                float(it.get("unit_price") or it.get("rate") or 0.0),
                float(it.get("taxable_amount") or it.get("taxable_value") or 0.0),
                float(it.get("total_amount") or 0.0)
            ])

    # Sheet 3: Validation Report
    ws_val = wb.create_sheet(title="Validation Report")
    val_headers = ["Bill Number", "Rule Status", "Errors / Checks"]
    ws_val.append(val_headers)
    for col in range(1, len(val_headers) + 1):
        ws_val.cell(row=1, column=col).font = header_font
        ws_val.cell(row=1, column=col).fill = header_fill

    for inv in invoices:
        val = inv.get("validation_json", {})
        passed = ", ".join(val.get("checks_passed", []))
        errors = "; ".join([e.get("message", "") for e in val.get("errors", [])])
        ws_val.append([
            inv.get("bill_number", ""),
            val.get("status", "valid"),
            errors if errors else (passed or "All checks passed")
        ])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
